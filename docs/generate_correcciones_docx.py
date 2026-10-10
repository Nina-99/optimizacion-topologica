"""Genera el documento Word de correcciones a partir del Markdown.

Lee `docs/Correcciones_Documento_Principal.md` y produce
`docs/Correcciones_Documento_Principal.docx` (formato académico, Times New Roman).

Uso:  python docs/generate_correcciones_docx.py
"""

import os
import re

from docx import Document
from docx.shared import Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE_DIR, "Correcciones_Documento_Principal.md")
OUT = os.path.join(BASE_DIR, "Correcciones_Documento_Principal.docx")

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")


def _add_runs(par, text):
    """Agrega texto a un párrafo manejando **negrita** y `code`."""
    pos = 0
    # tokens: **negrita** o `code`
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            r = par.add_run(m.group(1)); r.bold = True
        else:
            r = par.add_run(m.group(2)); r.font.name = "Consolas"; r.font.size = Pt(10)
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def _is_sep(row):
    return all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in row if c.strip() != "")


def _split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_text):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            i += 1
            continue

        # Bloque de código ```
        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"; r.font.size = Pt(10)
            continue

        # Tabla (| ... |)
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_row(lines[i])
                if not _is_sep(cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncols)
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(ncols):
                        txt = row[ci] if ci < len(row) else ""
                        cells[ci].text = ""
                        _add_runs(cells[ci].paragraphs[0], txt)
                        if ri == 0:
                            for run in cells[ci].paragraphs[0].runs:
                                run.bold = True
            continue

        # Encabezados
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=0); i += 1; continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1); i += 1; continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2); i += 1; continue

        # Separador ---
        if stripped == "---":
            i += 1; continue

        # Lista
        if stripped.startswith("- "):
            _add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
            i += 1; continue

        # Párrafo normal
        _add_runs(doc.add_paragraph(), stripped)
        i += 1

    return doc


def main():
    with open(SRC, encoding="utf-8") as f:
        md = f.read()
    doc = convert(md)
    doc.save(OUT)
    print(f"[OK] Documento de correcciones generado: {OUT}")


if __name__ == "__main__":
    main()
